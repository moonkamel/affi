"""Reporte docs/prospection/envoi.json dans docs/suivi-emailing.xlsx et génère le CSV
(formules au format Google Sheets français) à importer dans le Drive.
Usage : python3 docs/prospection/maj_suivi.py <sortie.csv>"""
import csv, datetime, json, sys
import openpyxl

XLSX = 'docs/suivi-emailing.xlsx'
data = json.load(open('docs/prospection/envoi.json'))
wb = openpyxl.load_workbook(XLSX)
ws = wb['Suivi']

def relance(d):
    r = d + datetime.timedelta(days=5)
    while r.weekday() >= 5:  # pas de relance le week-end
        r += datetime.timedelta(days=1)
    return r

for row in range(2, ws.max_row + 1):
    mail = ws.cell(row, 5).value
    info = data['envoyes'].get(mail)
    if not info:
        continue
    d = datetime.date.fromisoformat(info['date'])
    ws.cell(row, 10).value = d
    if info.get('relance'):
        ws.cell(row, 12).value = datetime.date.fromisoformat(info['relance'])
    statut = info.get('statut') or ('Relancé' if info.get('relance') else 'Envoyé')
    ws.cell(row, 9).value = statut
    if info.get('reponse'):
        ws.cell(row, 13).value = info['reponse']
    if info.get('lien'):
        ws.cell(row, 14).value = info['lien']
    if info.get('note'):
        ws.cell(row, 15).value = info['note']
    elif statut in ('Envoyé', 'Relancé'):
        ws.cell(row, 15).value = 'Relancé' if info.get('relance') else 'Relance prévue le ' + relance(d).strftime('%d/%m/%Y')
    else:
        ws.cell(row, 15).value = None
wb.save(XLSX)

# CSV pour Google Sheets (séparateur d'arguments « ; »)
def fr(v):
    if not isinstance(v, str) or not v.startswith('='):
        return v
    out, q = [], False
    for ch in v:
        if ch == '"':
            q = not q
        out.append(';' if ch == ',' and not q else ch)
    return ''.join(out)

N = 200
def rg(c): return f'${c}$2:${c}${N}'
stats = [('Tableau de bord', ''), ('Contacts avec adresse', f'=COUNTIF({rg("E")},"*@*")'),
         ('Brouillons prêts', f'=COUNTIF({rg("I")},"Brouillon prêt")'), ('Envoyés (tous suivis)', f'=COUNT({rg("J")})'),
         ('Relancés', f'=COUNTIF({rg("I")},"Relancé")'), ('Réponses reçues', f'=COUNTIF({rg("M")},"Oui")'),
         ('Liens obtenus', f'=COUNTIF({rg("I")},"Lien obtenu")'), ('Refus', f'=COUNTIF({rg("I")},"Refus")'),
         ('Adresses invalides', f'=COUNTIF({rg("I")},"Adresse invalide")'), ('Taux de réponse', '=IFERROR(R7/R5,0)'),
         ("Taux d'obtention de lien", '=IFERROR(R8/R5,0)'),
         ("Relances dues aujourd'hui", f'=COUNTIFS({rg("K")},"<="&TODAY(),{rg("I")},"Envoyé",{rg("L")},"")'),
         ('', ''), ('Par catégorie', 'Contacts', 'Envoyés', 'Liens')]
for cat in ['Média', 'Expert-comptable', 'Fédération', 'CCI', 'CMA', 'Réseau', 'Initiative', 'Réseau Entreprendre', 'Couveuse / CAE', 'Association']:
    r = len(stats) + 2
    stats.append((cat, f'=COUNTIF({rg("B")},Q{r})', f'=COUNTIFS({rg("B")},Q{r},{rg("J")},"<>")',
                  f'=COUNTIFS({rg("B")},Q{r},{rg("I")},"Lien obtenu")'))
rows = []
for r in range(1, ws.max_row + 1):
    if r > 1 and not ws.cell(r, 3).value:
        break
    vals = []
    for c in range(1, 16):
        v = ws.cell(r, c).value
        if c == 11 and r > 1:
            v = f'=IF(J{r}="","",WORKDAY(J{r}+4,1))'  # J+5, décalé au lundi si week-end
        if isinstance(v, datetime.datetime):
            v = v.strftime('%d/%m/%Y')
        elif isinstance(v, datetime.date):
            v = v.strftime('%d/%m/%Y')
        vals.append('' if v is None else v)
    rows.append(vals)
rows[0] += ['', 'Statistiques', '', '', '']
for i, s in enumerate(stats):
    s = list(s) + [''] * (4 - len(s))
    rows[i + 1] += [''] + s[:4]
with open(sys.argv[1], 'w', newline='') as f:
    csv.writer(f).writerows([[fr(v) for v in r] for r in rows])
print(len(rows) - 1, 'lignes')
